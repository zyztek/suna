import os
import io
import zipfile
import tempfile
import shutil
import asyncio
import subprocess
import re
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path
import mimetypes
import chardet

import PyPDF2
import docx
import openpyxl
import csv
import json
import yaml
import xml.etree.ElementTree as ET
from PIL import Image
import pytesseract

from utils.logger import logger
from services.supabase import DBConnection

class FileProcessor:
    """Handles file upload, content extraction, and processing for agent knowledge bases."""
    
    SUPPORTED_TEXT_EXTENSIONS = {
        '.txt', '.md', '.py', '.js', '.ts', '.html', '.css', '.json', '.yaml', '.yml',
        '.xml', '.csv', '.sql', '.sh', '.bat', '.ps1', '.dockerfile', '.gitignore',
        '.env', '.ini', '.cfg', '.conf', '.log', '.rst', '.toml', '.lock'
    }
    
    SUPPORTED_DOCUMENT_EXTENSIONS = {
        '.pdf', '.docx', '.xlsx', '.pptx'
    }
    
    SUPPORTED_IMAGE_EXTENSIONS = {
        '.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tiff', '.webp'
    }
    
    MAX_FILE_SIZE = 50 * 1024 * 1024
    MAX_ZIP_ENTRIES = 1000
    MAX_CONTENT_LENGTH = 100000
    
    def __init__(self):
        self.db = DBConnection()
    
    async def process_file_upload(
        self, 
        agent_id: str, 
        account_id: str, 
        file_content: bytes, 
        filename: str, 
        mime_type: str
    ) -> Dict[str, Any]:
        """Process a single uploaded file and extract its content."""
        try:
            file_size = len(file_content)
            if file_size > self.MAX_FILE_SIZE:
                raise ValueError(f"File too large: {file_size} bytes (max: {self.MAX_FILE_SIZE})")
            
            file_extension = Path(filename).suffix.lower()

            if file_extension == '.zip':
                return await self._process_zip_file(agent_id, account_id, file_content, filename)
            
            content = await self._extract_file_content(file_content, filename, mime_type)
            
            if not content or not content.strip():
                raise ValueError(f"No extractable content found in {filename}")
            
            client = await self.db.client
            
            entry_data = {
                'agent_id': agent_id,
                'account_id': account_id,
                'name': f"📄 {filename}",
                'description': f"Content extracted from uploaded file: {filename}",
                'content': content[:self.MAX_CONTENT_LENGTH],
                'source_type': 'file',
                'source_metadata': {
                    'filename': filename,
                    'mime_type': mime_type,
                    'file_size': file_size,
                    'extraction_method': self._get_extraction_method(file_extension, mime_type)
                },
                'file_size': file_size,
                'file_mime_type': mime_type,
                'usage_context': 'always',
                'is_active': True
            }
            
            result = await client.table('agent_knowledge_base_entries').insert(entry_data).execute()
            
            if not result.data:
                raise Exception("Failed to create knowledge base entry")
            
            return {
                'success': True,
                'entry_id': result.data[0]['entry_id'],
                'filename': filename,
                'content_length': len(content),
                'extraction_method': entry_data['source_metadata']['extraction_method']
            }
            
        except Exception as e:
            logger.error(f"Error processing file {filename}: {str(e)}")
            return {
                'success': False,
                'filename': filename,
                'error': str(e)
            }
    
    async def _process_zip_file(
        self, 
        agent_id: str, 
        account_id: str, 
        zip_content: bytes, 
        zip_filename: str
    ) -> Dict[str, Any]:
        """Extract and process all files from a ZIP archive."""
        
        try:
            client = await self.db.client
            
            zip_entry_data = {
                'agent_id': agent_id,
                'account_id': account_id,
                'name': f"📦 {zip_filename}",
                'description': f"ZIP archive: {zip_filename}",
                'content': f"ZIP archive containing multiple files. Extracted files will appear as separate entries.",
                'source_type': 'file',
                'source_metadata': {
                    'filename': zip_filename,
                    'mime_type': 'application/zip',
                    'file_size': len(zip_content),
                    'is_zip_container': True
                },
                'file_size': len(zip_content),
                'file_mime_type': 'application/zip',
                'usage_context': 'always',
                'is_active': True
            }
            
            zip_result = await client.table('agent_knowledge_base_entries').insert(zip_entry_data).execute()
            zip_entry_id = zip_result.data[0]['entry_id']
            
            # Extract files from ZIP
            extracted_files = []
            failed_files = []
            
            with zipfile.ZipFile(io.BytesIO(zip_content), 'r') as zip_ref:
                file_list = zip_ref.namelist()
                
                if len(file_list) > self.MAX_ZIP_ENTRIES:
                    raise ValueError(f"ZIP contains too many files: {len(file_list)} (max: {self.MAX_ZIP_ENTRIES})")
                
                for file_path in file_list:
                    if file_path.endswith('/'):
                        continue
                    
                    try:
                        file_content = zip_ref.read(file_path)
                        filename = os.path.basename(file_path)
                        
                        if not filename:  # Skip if no filename
                            continue
                        
                        # Detect MIME type
                        mime_type, _ = mimetypes.guess_type(filename)
                        if not mime_type:
                            mime_type = 'application/octet-stream'
                        
                        # Extract content
                        content = await self._extract_file_content(file_content, filename, mime_type)
                        
                        if content and content.strip():
                            extracted_entry_data = {
                                'agent_id': agent_id,
                                'account_id': account_id,
                                'name': f"📄 {filename}",
                                'description': f"Extracted from {zip_filename}: {file_path}",
                                'content': content[:self.MAX_CONTENT_LENGTH],
                                'source_type': 'zip_extracted',
                                'source_metadata': {
                                    'filename': filename,
                                    'original_path': file_path,
                                    'zip_filename': zip_filename,
                                    'mime_type': mime_type,
                                    'file_size': len(file_content),
                                    'extraction_method': self._get_extraction_method(Path(filename).suffix.lower(), mime_type)
                                },
                                'file_size': len(file_content),
                                'file_mime_type': mime_type,
                                'extracted_from_zip_id': zip_entry_id,
                                'usage_context': 'always',
                                'is_active': True
                            }
                            
                            extracted_result = await client.table('agent_knowledge_base_entries').insert(extracted_entry_data).execute()
                            
                            extracted_files.append({
                                'filename': filename,
                                'path': file_path,
                                'entry_id': extracted_result.data[0]['entry_id'],
                                'content_length': len(content)
                            })
                        
                    except Exception as e:
                        logger.error(f"Error extracting {file_path} from ZIP: {str(e)}")
                        failed_files.append({
                            'filename': os.path.basename(file_path),
                            'path': file_path,
                            'error': str(e)
                        })
            
            return {
                'success': True,
                'zip_entry_id': zip_entry_id,
                'zip_filename': zip_filename,
                'extracted_files': extracted_files,
                'failed_files': failed_files,
                'total_extracted': len(extracted_files),
                'total_failed': len(failed_files)
            }
            
        except Exception as e:
            logger.error(f"Error processing ZIP file {zip_filename}: {str(e)}")
            return {
                'success': False,
                'zip_filename': zip_filename,
                'error': str(e)
            }
    
    async def process_git_repository(
        self, 
        agent_id: str, 
        account_id: str, 
        git_url: str,
        branch: str = 'main',
        include_patterns: List[str] = None,
        exclude_patterns: List[str] = None
    ) -> Dict[str, Any]:
        """Clone a Git repository and extract content from supported files."""
        
        if include_patterns is None:
            include_patterns = ['*.py', '*.js', '*.ts', '*.md', '*.txt', '*.json', '*.yaml', '*.yml']
        
        if exclude_patterns is None:
            exclude_patterns = ['node_modules/*', '.git/*', '*.pyc', '__pycache__/*', '.env', '*.log']
        
        temp_dir = None
        try:
            # Create temporary directory
            temp_dir = tempfile.mkdtemp()
            
            # Clone repository
            clone_cmd = ['git', 'clone', '--depth', '1', '--branch', branch, git_url, temp_dir]
            process = await asyncio.create_subprocess_exec(
                *clone_cmd,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE
            )
            stdout, stderr = await process.communicate()
            
            if process.returncode != 0:
                raise Exception(f"Git clone failed: {stderr.decode()}")
            
            # Create main repository entry
            client = await self.db.client
            
            repo_name = git_url.split('/')[-1].replace('.git', '')
            repo_entry_data = {
                'agent_id': agent_id,
                'account_id': account_id,
                'name': f"🔗 {repo_name}",
                'description': f"Git repository: {git_url} (branch: {branch})",
                'content': f"Git repository cloned from {git_url}. Individual files are processed as separate entries.",
                'source_type': 'git_repo',
                'source_metadata': {
                    'git_url': git_url,
                    'branch': branch,
                    'include_patterns': include_patterns,
                    'exclude_patterns': exclude_patterns
                },
                'usage_context': 'always',
                'is_active': True
            }
            
            repo_result = await client.table('agent_knowledge_base_entries').insert(repo_entry_data).execute()
            repo_entry_id = repo_result.data[0]['entry_id']
            
            # Process files in repository
            processed_files = []
            failed_files = []
            
            for root, dirs, files in os.walk(temp_dir):
                # Skip .git directory
                if '.git' in dirs:
                    dirs.remove('.git')
                
                for file in files:
                    file_path = os.path.join(root, file)
                    relative_path = os.path.relpath(file_path, temp_dir)
                    
                    # Check if file should be included
                    if not self._should_include_file(relative_path, include_patterns, exclude_patterns):
                        continue
                    
                    try:
                        with open(file_path, 'rb') as f:
                            file_content = f.read()
                        
                        if len(file_content) > self.MAX_FILE_SIZE:
                            continue  # Skip large files
                        
                        # Detect MIME type
                        mime_type, _ = mimetypes.guess_type(file)
                        if not mime_type:
                            mime_type = 'application/octet-stream'
                        
                        # Extract content
                        content = await self._extract_file_content(file_content, file, mime_type)
                        
                        if content and content.strip():
                            # Create entry for file
                            file_entry_data = {
                                'agent_id': agent_id,
                                'account_id': account_id,
                                'name': f"📄 {file}",
                                'description': f"From {repo_name}: {relative_path}",
                                'content': content[:self.MAX_CONTENT_LENGTH],
                                'source_type': 'git_repo',
                                'source_metadata': {
                                    'filename': file,
                                    'relative_path': relative_path,
                                    'git_url': git_url,
                                    'branch': branch,
                                    'repo_name': repo_name,
                                    'mime_type': mime_type,
                                    'file_size': len(file_content),
                                    'extraction_method': self._get_extraction_method(Path(file).suffix.lower(), mime_type)
                                },
                                'file_size': len(file_content),
                                'file_mime_type': mime_type,
                                'extracted_from_zip_id': repo_entry_id,  # Reuse this field for git repo reference
                                'usage_context': 'always',
                                'is_active': True
                            }
                            
                            file_result = await client.table('agent_knowledge_base_entries').insert(file_entry_data).execute()
                            
                            processed_files.append({
                                'filename': file,
                                'relative_path': relative_path,
                                'entry_id': file_result.data[0]['entry_id'],
                                'content_length': len(content)
                            })
                    
                    except Exception as e:
                        logger.error(f"Error processing {relative_path} from git repo: {str(e)}")
                        failed_files.append({
                            'filename': file,
                            'relative_path': relative_path,
                            'error': str(e)
                        })
            
            return {
                'success': True,
                'repo_entry_id': repo_entry_id,
                'repo_name': repo_name,
                'git_url': git_url,
                'branch': branch,
                'processed_files': processed_files,
                'failed_files': failed_files,
                'total_processed': len(processed_files),
                'total_failed': len(failed_files)
            }
            
        except Exception as e:
            logger.error(f"Error processing git repository {git_url}: {str(e)}")
            return {
                'success': False,
                'git_url': git_url,
                'error': str(e)
            }
        
        finally:
            # Clean up temporary directory
            if temp_dir and os.path.exists(temp_dir):
                shutil.rmtree(temp_dir, ignore_errors=True)
    
    async def _extract_file_content(self, file_content: bytes, filename: str, mime_type: str) -> str:
        """Extract text content from various file types."""
        file_extension = Path(filename).suffix.lower()
        
        try:
            # Text files
            if file_extension in self.SUPPORTED_TEXT_EXTENSIONS or mime_type.startswith('text/'):
                return self._extract_text_content(file_content)
            
            # PDF files
            elif file_extension == '.pdf':
                return self._extract_pdf_content(file_content)
            
            # Word documents
            elif file_extension == '.docx':
                return self._extract_docx_content(file_content)
            
            # Excel files
            elif file_extension == '.xlsx':
                return self._extract_xlsx_content(file_content)
            
            # Images (OCR)
            elif file_extension in self.SUPPORTED_IMAGE_EXTENSIONS:
                return self._extract_image_content(file_content)
            
            # JSON files
            elif file_extension == '.json':
                return self._extract_json_content(file_content)
            
            # YAML files
            elif file_extension in {'.yaml', '.yml'}:
                return self._extract_yaml_content(file_content)
            
            # XML files
            elif file_extension == '.xml':
                return self._extract_xml_content(file_content)
            
            # CSV files
            elif file_extension == '.csv':
                return self._extract_csv_content(file_content)
            
            else:
                # Try to extract as text if possible
                return self._extract_text_content(file_content)
        
        except Exception as e:
            logger.error(f"Error extracting content from {filename}: {str(e)}")
            return f"Error extracting content: {str(e)}"
    
    def _extract_text_content(self, file_content: bytes) -> str:
        """Extract content from text files with encoding detection."""

        detected = chardet.detect(file_content)
        encoding = detected.get('encoding', 'utf-8')
        
        try:
            raw_text = file_content.decode(encoding)
        except UnicodeDecodeError:
            raw_text = file_content.decode('utf-8', errors='replace')
        
        return self._sanitize_content(raw_text)
    
    def _extract_pdf_content(self, file_content: bytes) -> str:
        """Extract text from PDF files."""
        
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
        text_content = []
        
        for page in pdf_reader.pages:
            text_content.append(page.extract_text())
        
        raw_text = '\n\n'.join(text_content)
        return self._sanitize_content(raw_text)
    
    def _extract_docx_content(self, file_content: bytes) -> str:
        """Extract text from Word documents."""
        
        doc = docx.Document(io.BytesIO(file_content))
        text_content = []
        
        for paragraph in doc.paragraphs:
            text_content.append(paragraph.text)
        
        raw_text = '\n'.join(text_content)
        return self._sanitize_content(raw_text)
    
    def _extract_xlsx_content(self, file_content: bytes) -> str:
        """Extract text from Excel files."""
        
        workbook = openpyxl.load_workbook(io.BytesIO(file_content))
        text_content = []
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text_content.append(f"Sheet: {sheet_name}")
            
            for row in sheet.iter_rows(values_only=True):
                row_text = [str(cell) if cell is not None else '' for cell in row]
                if any(row_text): 
                    text_content.append('\t'.join(row_text))
        
        raw_text = '\n'.join(text_content)
        return self._sanitize_content(raw_text)
    
    def _extract_image_content(self, file_content: bytes) -> str:
        """Extract text from images using OCR."""
        
        try:
            image = Image.open(io.BytesIO(file_content))
            raw_text = pytesseract.image_to_string(image)
            return self._sanitize_content(raw_text)
        except Exception as e:
            return f"OCR extraction failed: {str(e)}"
    
    def _extract_json_content(self, file_content: bytes) -> str:
        """Extract and format JSON content."""
        
        text = self._extract_text_content(file_content)
        try:
            parsed = json.loads(text)
            formatted = json.dumps(parsed, indent=2)
            return self._sanitize_content(formatted)
        except json.JSONDecodeError:
            return self._sanitize_content(text)
    
    def _extract_yaml_content(self, file_content: bytes) -> str:
        """Extract and format YAML content."""
        
        text = self._extract_text_content(file_content)
        try:
            parsed = yaml.safe_load(text)
            formatted = yaml.dump(parsed, default_flow_style=False)
            return self._sanitize_content(formatted)
        except yaml.YAMLError:
            return self._sanitize_content(text)
    
    def _extract_xml_content(self, file_content: bytes) -> str:
        """Extract content from XML files."""
        
        try:
            root = ET.fromstring(file_content)
            xml_string = ET.tostring(root, encoding='unicode')
            return self._sanitize_content(xml_string)
        except ET.ParseError:
            return self._extract_text_content(file_content)
    
    def _extract_csv_content(self, file_content: bytes) -> str:
        """Extract and format CSV content."""
        
        text = self._extract_text_content(file_content)
        try:
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
            formatted = '\n'.join(['\t'.join(row) for row in rows])
            return self._sanitize_content(formatted)
        except Exception:
            return self._sanitize_content(text)
    
    def _sanitize_content(self, content: str) -> str:
        """Sanitize extracted content to remove problematic characters for PostgreSQL."""
        if not content:
            return content

        sanitized = ''.join(char for char in content if ord(char) >= 32 or char in '\n\r\t')

        sanitized = sanitized.replace('\x00', '')
        sanitized = sanitized.replace('\u0000', '')
        
        sanitized = sanitized.replace('\ufeff', '')
        
        sanitized = sanitized.replace('\r\n', '\n').replace('\r', '\n')

        sanitized = re.sub(r'\n{4,}', '\n\n\n', sanitized)

        return sanitized.strip()

    def _get_extraction_method(self, file_extension: str, mime_type: str) -> str:
        """Get the extraction method used for a file type."""
        
        if file_extension == '.pdf':
            return 'PyPDF2'
        elif file_extension == '.docx':
            return 'python-docx'
        elif file_extension == '.xlsx':
            return 'openpyxl'
        elif file_extension in self.SUPPORTED_IMAGE_EXTENSIONS:
            return 'pytesseract OCR'
        elif file_extension == '.json':
            return 'JSON parser'
        elif file_extension in {'.yaml', '.yml'}:
            return 'YAML parser'
        elif file_extension == '.xml':
            return 'XML parser'
        elif file_extension == '.csv':
            return 'CSV parser'
        else:
            return 'text encoding detection'
    
    def _should_include_file(self, file_path: str, include_patterns: List[str], exclude_patterns: List[str]) -> bool:
        """Check if a file should be included based on patterns."""
        
        import fnmatch
        
        for pattern in exclude_patterns:
            if fnmatch.fnmatch(file_path, pattern):
                return False
        
        for pattern in include_patterns:
            if fnmatch.fnmatch(file_path, pattern):
                return True
        
        return False 